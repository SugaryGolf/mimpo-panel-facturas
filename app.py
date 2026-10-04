from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
import pyodbc
import xml.etree.ElementTree as ET
import pandas as pd
from io import BytesIO
import re
# Se agregó Alignment para centrar las celdas
from openpyxl.styles import PatternFill, Font, Alignment 

app = Flask(__name__)
CORS(app) 

DB_NAME = 'PanelFacturas'

def obtener_conexion():
    drivers_disponibles = pyodbc.drivers()
    driver_elegido = 'SQL Server'
    for driver in ['ODBC Driver 18 for SQL Server', 'ODBC Driver 17 for SQL Server', 'SQL Server Native Client 11.0', 'SQL Server']:
        if driver in drivers_disponibles:
            driver_elegido = driver
            break
    extra_param = ';TrustServerCertificate=yes' if '18' in driver_elegido else ''
    servidores_a_probar = [r'(localdb)\LOCALDB#50C632A7', r'(localdb)\MSSQLLocalDB', r'.\SQLEXPRESS', r'.']
    
    ultimo_error = None
    for server in servidores_a_probar:
        try:
            return pyodbc.connect(f'DRIVER={{{driver_elegido}}};SERVER={server};DATABASE={DB_NAME};Trusted_Connection=yes{extra_param};', timeout=3)
        except Exception as e:
            ultimo_error = e
    raise ultimo_error

@app.route('/api/dashboard', methods=['GET'])
def obtener_dashboard():
    try:
        conn = obtener_conexion()
        cursor = conn.cursor()
        
        # 1. KPIs Generales
        cursor.execute("SELECT COUNT(Id), SUM(Total) FROM Facturas")
        resultado = cursor.fetchone()
        
        # 2. Historial de Facturas (AHORA EXTRAEMOS FOLIO, SUBTOTAL E IVA)
        cursor.execute("""
            SELECT Id, UUID, Nombre_Emisor, Fecha, Total, Folio, Subtotal, IVA_Trasladado 
            FROM Facturas 
            ORDER BY Total DESC
        """)
        facturas_db = cursor.fetchall()
        
        # Extraemos los equipos para cruzar las Series y Páginas con cada factura
        cursor.execute("SELECT Factura_Id, Numero_Serie, Paginas_Impresas FROM Equipos")
        equipos_db = cursor.fetchall()
        
        equipos_dict = {}
        for eq in equipos_db:
            f_id = eq.Factura_Id
            if f_id not in equipos_dict:
                equipos_dict[f_id] = {'series': set(), 'paginas': 0}
            if eq.Numero_Serie and eq.Numero_Serie != 'SIN SERIE':
                # Si vienen varias series pegadas con '|', las separamos limpiamente
                for s in eq.Numero_Serie.split('|'):
                    equipos_dict[f_id]['series'].add(s.strip())
            equipos_dict[f_id]['paginas'] += (eq.Paginas_Impresas or 0)

        # Armamos la lista final enviando TODO a React
        facturas = []
        for f in facturas_db:
            eq_info = equipos_dict.get(f.Id, {'series': set(), 'paginas': 0})
            series_str = " | ".join(sorted(list(eq_info['series']))) if eq_info['series'] else "N/A"
            
            facturas.append({
                "uuid": f.UUID, 
                "emisor": f.Nombre_Emisor, 
                "fecha": f.Fecha.strftime('%Y-%m-%d') if f.Fecha else '', 
                "total": float(f.Total) if f.Total else 0.0,
                "folio": f.Folio if f.Folio else "S/N",
                "subtotal": float(f.Subtotal) if f.Subtotal else 0.0,
                "iva_trasladado": float(f.IVA_Trasladado) if f.IVA_Trasladado else 0.0,
                "serie": series_str,
                "paginas": eq_info['paginas']
            })
        
        # 3. Datos para Gráfica de Comparativa Mes a Mes
        cursor.execute("""
            SELECT SUBSTRING(CONVERT(VARCHAR, Fecha, 120), 1, 7) as Mes, SUM(Total) 
            FROM Facturas 
            WHERE Fecha IS NOT NULL
            GROUP BY SUBSTRING(CONVERT(VARCHAR, Fecha, 120), 1, 7)
            ORDER BY Mes
        """)
        meses = [{"mes": row[0], "total": float(row[1])} for row in cursor.fetchall()]

        # 4. Datos para Gráfica de Gasto por Emisor
        cursor.execute("""
            SELECT TOP 5 Nombre_Emisor, SUM(Total) 
            FROM Facturas 
            GROUP BY Nombre_Emisor
            ORDER BY SUM(Total) DESC
        """)
        emisores = [{"name": row[0][:20] + "..." if len(row[0]) > 20 else row[0], "value": float(row[1])} for row in cursor.fetchall()]

        # 5. Equipos con más páginas impresas (Ranking)
        cursor.execute("""
            SELECT Numero_Serie, SUM(Paginas_Impresas) as Total_Paginas
            FROM Equipos
            WHERE Numero_Serie != 'SIN SERIE'
            GROUP BY Numero_Serie
            HAVING SUM(Paginas_Impresas) > 0
            ORDER BY Total_Paginas DESC
        """)
        equipos_ranking = [{"serie": row[0], "paginas": int(row[1])} for row in cursor.fetchall()]

        conn.close()
        
        return jsonify({
            "kpis": {"totalFacturas": resultado[0] or 0, "montoTotal": float(resultado[1]) if resultado[1] else 0.0},
            "facturas": facturas,
            "graficas": {
                "mesAMes": meses,
                "emisores": emisores,
                "equipos": equipos_ranking
            }
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/upload', methods=['POST'])
def subir_factura():
    if 'file' not in request.files: return jsonify({"error": "No se enviaron archivos"}), 400
    
    # Recibimos una lista de archivos gracias al "multiple" de React
    archivos = request.files.getlist('file')
    exitos = 0
    errores = 0
    
    conn = obtener_conexion()
    cursor = conn.cursor()

    for archivo in archivos:
        try:
            root = ET.parse(archivo).getroot()
            fecha_completa = root.attrib.get('Fecha', '')
            fecha, hora = fecha_completa.split('T') if 'T' in fecha_completa else (fecha_completa, '00:00:00')
            folio = root.attrib.get('Folio', '')
            total, subtotal, descuento = root.attrib.get('Total', 0), root.attrib.get('Subtotal', 0), root.attrib.get('Descuento', 0)
            moneda = root.attrib.get('Moneda', 'MXN')

            emisor, receptor, conceptos, uuid, iva_trasladado = {}, {}, [], "NO_UUID", 0.0
            series, paginas_array = [], []
            
            for elem in root.iter():
                tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
                if tag == 'Emisor': emisor = elem.attrib
                elif tag == 'Receptor': receptor = elem.attrib
                elif tag == 'Concepto': conceptos.append(elem.attrib)
                elif tag == 'TimbreFiscalDigital': uuid = elem.attrib.get('UUID', 'NO_UUID')
                elif tag == 'Impuestos' and elem.attrib.get('TotalImpuestosTrasladados'): 
                    iva_trasladado = elem.attrib.get('TotalImpuestosTrasladados')
                
                if tag == 'VolumenImp':
                    try: paginas_array.append(int(float(elem.attrib.get('Volumen', '0').replace(',', ''))))
                    except: pass
                
                for k, v in elem.attrib.items():
                    if re.match(r'^(serie|serial|ns|sn)$', k, re.IGNORECASE) and v.strip().upper() not in series:
                        series.append(v.strip().upper())

            xml_string = ET.tostring(root, encoding='utf8').decode('utf8')
            matches = re.findall(r'(?:\bS\s*/\s*N\b|\bN\s*/\s*S\b|\bSN\b|\bNS\b|\bSERIAL\b|\bSERIE\b)(?:\s+(?:DEL|DE|EQUIPO|NO|NUM|N[UÚ]MERO)\b\.?)*\s*[:#.\-]*\s*(?=[A-Za-z0-9_-]*\d)([A-Za-z0-9][A-Za-z0-9_-]*)', xml_string, re.IGNORECASE)
            for m in matches:
                if m.strip().upper() not in series: series.append(m.strip().upper())

            if uuid == "NO_UUID":
                errores += 1
                continue

            # Verificamos duplicados
            if cursor.execute("SELECT Id FROM Facturas WHERE UUID = ?", uuid).fetchone():
                errores += 1
                continue

            cursor.execute("INSERT INTO Facturas (UUID, Fecha, Hora, Folio, Nombre_Emisor, Nombre_Receptor, Subtotal, Descuento, IVA_Trasladado, Total, Moneda) OUTPUT INSERTED.Id VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", 
                           (uuid, fecha, hora, folio, emisor.get('Nombre', ''), receptor.get('Nombre', ''), subtotal, descuento, iva_trasladado, total, moneda))
            factura_id = cursor.fetchone()[0]

            for c in conceptos:
                cursor.execute("INSERT INTO Conceptos (Factura_Id, ClaveProdServ, Descripcion, ClaveUnidad, Unidad, Cantidad, ValorUnitario, Importe) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                               (factura_id, c.get('ClaveProdServ',''), c.get('Descripcion',''), c.get('ClaveUnidad',''), c.get('Unidad',''), c.get('Cantidad',0), c.get('ValorUnitario',0), c.get('Importe',0)))
            
            series_str = " | ".join(series) if series else "SIN SERIE"
            if paginas_array:
                for p in paginas_array:
                    cursor.execute("INSERT INTO Equipos (Factura_Id, Numero_Serie, Paginas_Impresas) VALUES (?, ?, ?)", (factura_id, series_str, p))
            else:
                cursor.execute("INSERT INTO Equipos (Factura_Id, Numero_Serie, Paginas_Impresas) VALUES (?, ?, ?)", (factura_id, series_str, 0))

            exitos += 1
        except Exception:
            errores += 1

    conn.commit()
    conn.close()
    
    res_msg = f"¡Proceso completado! Se guardaron {exitos} facturas nuevas."
    if errores > 0:
        res_msg += f" ({errores} ignoradas/duplicadas)."
        
    return jsonify({"mensaje": res_msg})

@app.route('/api/export', methods=['GET'])
def exportar_excel():
    try:
        conn = obtener_conexion()
        df_f = pd.read_sql("SELECT * FROM Facturas", conn)
        df_c = pd.read_sql("SELECT Factura_Id, ClaveProdServ, Descripcion, ClaveUnidad, Unidad, Cantidad FROM Conceptos", conn)
        df_e = pd.read_sql("SELECT Factura_Id, Numero_Serie, Paginas_Impresas FROM Equipos", conn)
        conn.close()

        if df_f.empty: return jsonify({"error": "No hay datos para exportar"}), 404

        df_c_agg = df_c.groupby('Factura_Id').agg({
            'ClaveProdServ': lambda x: ' | '.join(x.dropna().astype(str)), 
            'Descripcion': lambda x: ' | '.join(x.dropna().astype(str)), 
            'Cantidad': 'sum', 
            'ClaveUnidad': lambda x: ' | '.join(pd.Series(x).dropna().astype(str).unique()), 
            'Unidad': lambda x: ' | '.join(pd.Series(x).dropna().astype(str).unique())
        }).reset_index()

        df_e_agg = df_e.groupby('Factura_Id').agg({
            'Numero_Serie': lambda x: ' | '.join(pd.Series(x).dropna().astype(str).unique()),
            'Paginas_Impresas': 'sum'
        }).reset_index()

        df_master = df_f.merge(df_c_agg, left_on='Id', right_on='Factura_Id', how='left').merge(df_e_agg, left_on='Id', right_on='Factura_Id', how='left')

        # ORDENAR DE MONTO MAYOR A MENOR
        df_master['Total'] = pd.to_numeric(df_master['Total'], errors='coerce').fillna(0)
        df_master = df_master.sort_values(by='Total', ascending=False).reset_index(drop=True)

        df_facturas = pd.DataFrame()
        df_facturas['Posición'] = range(1, len(df_master) + 1)
        df_facturas['UUID'] = df_master['UUID']
        df_facturas['Fecha'] = df_master['Fecha']
        df_facturas['Hora'] = df_master['Hora']
        df_facturas['Folio'] = df_master['Folio']
        df_facturas['Nombre Emisor'] = df_master['Nombre_Emisor']
        df_facturas['Nombre Receptor'] = df_master['Nombre_Receptor']
        df_facturas['Producto'] = df_master['ClaveProdServ']
        df_facturas['Descripción'] = df_master['Descripcion']
        df_facturas['Número de Serie'] = df_master['Numero_Serie']
        df_facturas['Total Páginas Impresas'] = df_master['Paginas_Impresas']
        df_facturas['Cantidad Total'] = df_master['Cantidad']
        df_facturas['Claves Unidad'] = df_master['ClaveUnidad']
        df_facturas['Unidades'] = df_master['Unidad']
        df_facturas['SubTotal'] = df_master['Subtotal']
        df_facturas['Descuento'] = df_master['Descuento']
        df_facturas['IVA Trasladado'] = df_master['IVA_Trasladado']
        df_facturas['Total Factura'] = df_master['Total']
        df_facturas['Moneda'] = df_master['Moneda']

        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df_facturas.to_excel(writer, index=False, sheet_name='Detalle_Facturas')
            
            # FORMATO: CENTRADO Y CABECERAS ROJAS
            header_fill = PatternFill(start_color='A6192E', end_color='A6192E', fill_type='solid')
            header_font = Font(color='FFFFFF', bold=True)
            center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

            ws = writer.book['Detalle_Facturas']
            
            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = center_alignment
                    if cell.row == 1:
                        cell.fill = header_fill
                        cell.font = header_font

            for col in ws.columns:
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = 25 # Ancho estándar amplio
                
        output.seek(0)
        return send_file(output, download_name="Reporte_Facturas_Maestro.xlsx", as_attachment=True)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True, port=5000)